#!/usr/bin/env python3
"""
Slackスレッド分析レポートのPowerPointスライドを生成する
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


def add_title_slide(prs, title, subtitle=""):
    """タイトルスライドを追加"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # タイトル
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    # サブタイトル
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(9), Inches(1))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(24)
        p.alignment = PP_ALIGN.CENTER

    return slide


def add_content_slide(prs, title, content_lines):
    """コンテンツスライドを追加"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # タイトル
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True

    # 区切り線
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5), Inches(1.1), Inches(9), Inches(0.03)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 112, 192)
    line.line.fill.background()

    # コンテンツ
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True

    for i, line_text in enumerate(content_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = line_text
        p.font.size = Pt(18)
        p.space_after = Pt(12)

    return slide


def add_table_slide(prs, title, headers, rows):
    """テーブル付きスライドを追加"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)

    # タイトル
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True

    # 区切り線
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.5), Inches(1.1), Inches(9), Inches(0.03)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0, 112, 192)
    line.line.fill.background()

    # テーブル
    num_rows = len(rows) + 1
    num_cols = len(headers)

    table = slide.shapes.add_table(
        num_rows, num_cols,
        Inches(0.5), Inches(1.5),
        Inches(9), Inches(0.5 * num_rows)
    ).table

    # ヘッダー
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0, 112, 192)
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)

    # データ行
    for row_idx, row in enumerate(rows):
        for col_idx, cell_text in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(cell_text)
            cell.text_frame.paragraphs[0].font.size = Pt(12)

    return slide


def create_presentation():
    """メインのプレゼンテーション作成"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # スライド1: タイトル
    add_title_slide(
        prs,
        "Slackスレッド分析レポート",
        "脊椎番号付け（Spine_Num D2）プロジェクト\n評価とToDo整理"
    )

    # スライド2: エグゼクティブサマリー
    add_content_slide(prs, "エグゼクティブサマリー", [
        "プロジェクト目的:",
        "　医療画像（CT/MRI）から脊椎番号を自動で付与するシステムの開発",
        "",
        "現在の状況:",
        "　✅ GitHub初期整備: 完了",
        "　⬜ Clean inference設計: 次のステップ",
        "　⬜ 軸判定自動化: 未対応（最重要課題）",
        "　⬜ DICOM locator出力: 完全未着手",
        "",
        "結論:",
        "　現行本線（TotalSegmentator + ConvNeXtV2）は妥当",
        "　次のアクション: Clean single-case inference entry point設計に進む",
    ])

    # スライド3: プロジェクト背景
    add_content_slide(prs, "プロジェクト背景", [
        "解決したい課題:",
        "　・脊椎は頸椎(C1-C7)、胸椎(T1-T12)、腰椎(L1-L5)等で構成",
        "　・手動での番号付けは時間がかかり、ミスも発生しやすい",
        "",
        "目指すワークフロー:",
        "　CT/MRI撮影 → AI自動番号付け → 診断・報告",
        "",
        "最終目標:",
        "　・DICOM locator形式での出力（スカウト画像として動作）",
        "　・臨床ワークフローへの統合",
    ])

    # スライド4: 現行アーキテクチャ
    add_content_slide(prs, "現行アーキテクチャ", [
        "パイプライン構成:",
        "",
        "　CT/MRI画像 (DICOM/NIfTI)",
        "　　　↓",
        "　TotalSegmentator (骨格抽出・104クラス対応)",
        "　　　↓",
        "　矢状断RaySum生成 (3D→2D投影) ★課題: 軸判定",
        "　　　↓",
        "　ConvNeXtV2 (26チャンネル特徴抽出)",
        "　　　↓",
        "　ヒートマップ回帰 (各脊椎位置を推定)",
        "　　　↓",
        "　脊椎番号出力 (C1非表示/C1除外評価)",
    ])

    # スライド5: ツール妥当性評価
    add_table_slide(
        prs,
        "ツール妥当性評価",
        ["ツール", "用途", "評価", "備考"],
        [
            ["TotalSegmentator", "3Dセグメンテーション", "✅ 妥当", "医療画像標準、104クラス対応"],
            ["ConvNeXtV2", "特徴抽出", "✅ 妥当", "SOTA CNN、ImageNet事前学習"],
            ["ヒートマップ回帰", "位置推定", "✅ 妥当", "landmark検出に実績あり"],
            ["Skellytour", "骨格検出", "⚠️ 要検証", "正常変異で不安定（探索段階）"],
        ]
    )

    # スライド6: 未解決課題
    add_table_slide(
        prs,
        "未解決課題一覧",
        ["#", "課題", "重要度", "状態", "詳細"],
        [
            ["1", "軸判定・向き判定の自動化", "🔴 高", "未対応", "TotalSegmentator→RaySum時に方向不明"],
            ["2", "RaySum投影方向の不整合", "🔴 高", "部分対応?", "元データでデタラメになる"],
            ["3", "DICOM locator形式出力", "🟡 中", "完全未着手", "スカウト同様の形式化"],
            ["4", "Skellytour正常変異対応", "🟡 中", "評価中", "highモデル未テスト"],
        ]
    )

    # スライド7: 補完すべき技術要素
    add_content_slide(prs, "補完すべき技術要素", [
        "1. DICOM方向情報パーサー",
        "　　├ ImageOrientationPatient解析",
        "　　└ 自動軸判定",
        "",
        "2. 品質管理モジュール",
        "　　├ 異常検知（解剖学的整合性チェック）",
        "　　└ 信頼度スコア出力",
        "",
        "3. DICOM SC/Locator生成モジュール",
        "　　├ pydicom使用",
        "　　└ locator情報埋め込み",
    ])

    # スライド8: 進捗状況
    add_content_slide(prs, "現在の進捗状況", [
        "[A] GitHub初期整備　　　　████████████████████ 100% ✅完了",
        "　　├ 仕様・参照コード整理",
        "　　├ docs/model_registry.md方針決定",
        "　　└ checkpoint管理方針決定",
        "",
        "[B] Clean inference設計　░░░░░░░░░░░░░░░░░░░░   0% ⬜次のステップ",
        "",
        "[C] 軸判定自動化　　　　　░░░░░░░░░░░░░░░░░░░░   0% ⬜未着手",
        "",
        "[D] DICOM locator出力　　░░░░░░░░░░░░░░░░░░░░   0% ⬜未着手",
        "",
        "[E] Skellytour評価　　　　████████░░░░░░░░░░░░  40% 🔄進行中",
    ])

    # スライド9: ToDo整理
    add_content_slide(prs, "ToDo整理（優先順位付き）", [
        "Phase 1: 基盤整備（即時）",
        "　[ ] 既存コードのインベントリ作成",
        "　[ ] 開発環境セットアップ文書化",
        "　[ ] checkpointファイルの管理確認",
        "",
        "Phase 2: 推論パイプライン（短期）⭐最優先",
        "　[ ] Clean single-case inference entry point実装",
        "　[ ] 軸判定・向き判定の自動化",
        "　[ ] RaySum投影方向の標準化",
        "",
        "Phase 3: 出力形式（中期）",
        "　[ ] DICOM locator形式出力の設計・実装",
    ])

    # スライド10: 開発方針
    add_content_slide(prs, "開発方針", [
        "1. 現行本線を優先開発",
        "　　TotalSegmentator + ConvNeXtV2",
        "　　→ 実績あり、安定性高い",
        "",
        "2. Skellytourは並行評価（探索ブランチ）",
        "　　highモデルで追加検証",
        "　　→ 良好なら統合、不良なら棄却",
        "",
        "3. 軸判定問題は最優先で解決",
        "　　DICOM方向情報活用を第一候補",
        "　　→ 2D分類はフォールバック",
    ])

    # スライド11: 次のステップ
    add_content_slide(prs, "次のステップ", [
        "具体的アクション: Clean single-case inference",
        "",
        "入力:",
        "　・DICOM/NIfTIファイルパス",
        "　・(オプション) 設定JSON",
        "",
        "出力:",
        "　・脊椎番号付きJSON",
        "　・DICOM locator (将来)",
        "",
        "含めるべきロジック:",
        "　・DICOM方向情報の解析",
        "　・軸判定の自動化",
        "　・エラーハンドリング",
    ])

    # スライド12: 確認事項
    add_content_slide(prs, "確認事項", [
        "Slackスレッドからの不明点:",
        "",
        "1. 2D分類モデルでRaySum投影方向対応は完了しているか？",
        "",
        "2. C1非表示/C1除外評価の詳細仕様は？",
        "",
        "3. テスト用サンプルデータの入手方法は？",
        "",
        "4. checkpointファイルの共有ストレージパスは？",
    ])

    return prs


if __name__ == "__main__":
    prs = create_presentation()
    output_path = "/home/user/SBA/slack_thread_analysis_slides.pptx"
    prs.save(output_path)
    print(f"PowerPoint saved to: {output_path}")
